"""Analyze, validate, visualize, and integrate empirical honeybee datasets."""

from __future__ import annotations

import json
import sys
import zipfile
from csv import DictReader
from io import BytesIO, TextIOWrapper
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.io import loadmat

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from beestack import BeeStackConfig, config_from_mapping, run_simulation, zero_observation
from beestack.brain import (
    AntennalMovementSummary,
    AtlasDownloadRecord,
    AtlasInventory,
    BeeBrainAnatomySummary,
    BeeBrainEndToEndReport,
    EmpiricalCalciumDataset,
    EmpiricalOdorResponsePanel,
    EmpiricalTemplateBank,
    EmpiricalWaggleFollowerDataset,
    NeuropilAbbreviation,
    activity_summary_from_components,
    all_checks_passed,
    analyze_odor_panel,
    antennal_vibration_from_movement,
    atlas_download_records_from_dicts,
    atlas_inventories_from_directory,
    bee_brain_data_completeness_panel,
    build_empirical_template_bank,
    empirical_brain_datasets,
    load_neuropil_abbreviations,
    parse_paoli_matlab_payload,
    parse_tabular_odor_response_rows,
    process_observation,
    response_templates_from_calcium_dataset,
    summarize_anatomy,
    summarize_antennal_movement_rows,
    summarize_calcium_trials,
    summarize_waggle_follower_rows,
    template_alignment_matrix,
    validate_calcium_dataset,
    waggle_vibration_from_followers,
)
from beestack.visualization import generate_empirical_figures
from signpost_project_tree import write_project_readiness_review, write_signposts


def load_config() -> BeeStackConfig:
    import yaml

    config_path = PROJECT_ROOT / "manuscript" / "config.yaml"
    if not config_path.exists():
        return BeeStackConfig()
    payload = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    return config_from_mapping(payload.get("beestack", payload))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def load_empirical_panels(source_dir: Path) -> tuple[EmpiricalOdorResponsePanel, ...]:
    panels: list[EmpiricalOdorResponsePanel] = []
    for zip_path in sorted(source_dir.glob("*/*.zip")):
        dataset_id = zip_path.parent.name
        with zipfile.ZipFile(zip_path) as archive:
            for name, data in _iter_xlsx_members(archive):
                panels.extend(_panels_from_workbook(dataset_id, name, data))
    for workbook_path in sorted(source_dir.glob("*/files/**/*.xlsx")):
        dataset_id = _dataset_id_from_path(source_dir, workbook_path)
        panels.extend(
            _panels_from_workbook(dataset_id, workbook_path.name, workbook_path.read_bytes())
        )
    for workbook_path in sorted(source_dir.glob("*/*.xlsx")):
        dataset_id = workbook_path.parent.name
        panels.extend(
            _panels_from_workbook(dataset_id, workbook_path.name, workbook_path.read_bytes())
        )
    return tuple(panels)


def load_antennal_movement_summaries(source_dir: Path) -> tuple[AntennalMovementSummary, ...]:
    """Stream real Jernigan antennal CSV files into compact movement summaries."""

    summaries: list[AntennalMovementSummary] = []
    for zip_path in sorted(source_dir.glob("*/*.zip")):
        dataset_id = zip_path.parent.name
        with zipfile.ZipFile(zip_path) as archive:
            summaries.extend(_antennal_summaries_from_archive(dataset_id, archive))
    return tuple(summaries)


def load_waggle_follower_dataset(source_dir: Path) -> EmpiricalWaggleFollowerDataset | None:
    """Load Figshare Hadjitofi-Webb waggle-following CSV files when local."""

    dataset_id = "figshare-hadjitofi-2024-waggle-following"
    dataset_dir = source_dir / dataset_id
    feature_rows: list[dict[str, Any]] = []
    binned_rows: list[dict[str, Any]] = []
    model_error_rows: list[dict[str, Any]] = []
    straightness_rows: list[dict[str, Any]] = []
    source_files: list[str] = []
    for csv_path in sorted(dataset_dir.glob("files/**/*.csv")) + sorted(dataset_dir.glob("*.csv")):
        rows = _csv_rows(csv_path)
        if not rows:
            continue
        source_files.append(str(csv_path))
        lowered = csv_path.name.lower()
        stamped = [dict(row, _source_file=csv_path.name) for row in rows]
        if "features" in lowered and {"bee_id", "angle_to_dancer_deg"} <= set(rows[0]):
            feature_rows.extend(stamped)
        elif "binned" in lowered and "errors" not in lowered:
            binned_rows.extend(stamped)
        elif "errors" in lowered:
            model_error_rows.extend(stamped)
        elif "straightness" in lowered:
            straightness_rows.extend(stamped)
    if not feature_rows:
        return None
    parsed = summarize_waggle_follower_rows(
        feature_rows,
        dataset_id=dataset_id,
        binned_rows=binned_rows,
        model_error_rows=model_error_rows,
        straightness_rows=straightness_rows,
    )
    return EmpiricalWaggleFollowerDataset(
        dataset_id=parsed.dataset_id,
        tracks=parsed.tracks,
        summary=parsed.summary,
        source_files=tuple(source_files),
    )


def load_calcium_datasets(source_dir: Path) -> tuple[EmpiricalCalciumDataset, ...]:
    """Load Paoli-style MATLAB calcium datasets from archives or direct files."""

    datasets: list[EmpiricalCalciumDataset] = []
    for zip_path in sorted(source_dir.glob("*/*.zip")):
        dataset_id = zip_path.parent.name
        with zipfile.ZipFile(zip_path) as archive:
            datasets.extend(_calcium_datasets_from_archive(dataset_id, archive))
    for mat_path in sorted(source_dir.glob("*/files/**/*.mat")):
        dataset = _calcium_dataset_from_mat_bytes(mat_path.read_bytes())
        if dataset is not None:
            datasets.append(dataset)
    for zip_path in sorted(source_dir.glob("*/files/**/*.mat.zip")):
        with zipfile.ZipFile(zip_path) as archive:
            datasets.extend(
                _calcium_datasets_from_archive(
                    _dataset_id_from_path(source_dir, zip_path), archive, prefix=f"{zip_path.name}/"
                )
            )
    return tuple(datasets)


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in DictReader(handle)]


def workbook_dataframe_summaries(source_dir: Path) -> list[dict[str, int | float | str]]:
    """Summarize workbook sheets through pandas dataframes for data QA."""

    summaries: list[dict[str, int | float | str]] = []
    for zip_path in sorted(source_dir.glob("*/*.zip")):
        dataset_id = zip_path.parent.name
        with zipfile.ZipFile(zip_path) as archive:
            for name, data in _iter_xlsx_members(archive):
                summaries.extend(_dataframe_summaries_from_workbook(dataset_id, name, data))
    for workbook_path in sorted(source_dir.glob("*/files/**/*.xlsx")):
        dataset_id = _dataset_id_from_path(source_dir, workbook_path)
        summaries.extend(
            _dataframe_summaries_from_workbook(
                dataset_id, workbook_path.name, workbook_path.read_bytes()
            )
        )
    for workbook_path in sorted(source_dir.glob("*/*.xlsx")):
        dataset_id = workbook_path.parent.name
        summaries.extend(
            _dataframe_summaries_from_workbook(
                dataset_id, workbook_path.name, workbook_path.read_bytes()
            )
        )
    return summaries


def load_anatomy(
    source_dir: Path,
) -> tuple[
    BeeBrainAnatomySummary,
    tuple[AtlasInventory, ...],
    tuple[NeuropilAbbreviation, ...],
    tuple[AtlasDownloadRecord, ...],
]:
    """Load downloaded Honeybee Standard Brain inventories and abbreviations."""

    rows = _json_rows(source_dir / "anatomy_downloads.json")
    records = atlas_download_records_from_dicts(tuple(rows))
    atlas_dir = source_dir / "virtual-honeybee-standard-brain"
    inventories = atlas_inventories_from_directory(atlas_dir) if atlas_dir.exists() else ()
    abbreviations = load_neuropil_abbreviations(atlas_dir / "neuropil_abbreviations.html")
    summary = summarize_anatomy(records, inventories, abbreviations)
    return summary, inventories, abbreviations, records


def _antennal_summaries_from_archive(
    dataset_id: str,
    archive: zipfile.ZipFile,
    prefix: str = "",
) -> list[AntennalMovementSummary]:
    summaries: list[AntennalMovementSummary] = []
    required_columns = {"Bee", "Plume", "Frame", "Odor_detector", "LeftTheta", "RightTheta"}
    for info in archive.infolist():
        if info.is_dir():
            continue
        member_name = f"{prefix}{info.filename}"
        lowered = info.filename.lower()
        if lowered.endswith(".csv"):
            with (
                archive.open(info) as raw,
                TextIOWrapper(raw, encoding="utf-8-sig", newline="") as text,
            ):
                rows = DictReader(text)
                fieldnames = set(rows.fieldnames or ())
                if required_columns <= fieldnames:
                    summaries.append(summarize_antennal_movement_rows(rows, dataset_id=dataset_id))
        elif lowered.endswith(".zip"):
            with zipfile.ZipFile(BytesIO(archive.read(info))) as nested:
                summaries.extend(
                    _antennal_summaries_from_archive(dataset_id, nested, prefix=f"{member_name}/")
                )
    return summaries


def _calcium_datasets_from_archive(
    dataset_id: str,
    archive: zipfile.ZipFile,
    prefix: str = "",
) -> list[EmpiricalCalciumDataset]:
    datasets: list[EmpiricalCalciumDataset] = []
    for info in archive.infolist():
        if info.is_dir():
            continue
        member_name = f"{prefix}{info.filename}"
        lowered = info.filename.lower()
        if lowered.endswith(".mat"):
            dataset = _calcium_dataset_from_mat_bytes(archive.read(info), member_name)
            if dataset is not None:
                datasets.append(dataset)
        elif lowered.endswith(".zip"):
            with zipfile.ZipFile(BytesIO(archive.read(info))) as nested:
                datasets.extend(
                    _calcium_datasets_from_archive(dataset_id, nested, prefix=f"{member_name}/")
                )
    return datasets


def _calcium_dataset_from_mat_bytes(
    data: bytes, source_name: str = "matlab_payload"
) -> EmpiricalCalciumDataset | None:
    try:
        payload = loadmat(BytesIO(data), squeeze_me=True, struct_as_record=False)
    except (NotImplementedError, ValueError):
        payload = _load_hdf5_mat_payload(data)
    except Exception:
        return None
    try:
        dataset = parse_paoli_matlab_payload(payload)
    except (TypeError, ValueError, KeyError):
        return None
    return EmpiricalCalciumDataset(
        dataset_id=dataset.dataset_id,
        traces=dataset.traces,
        acquisition_hz=dataset.acquisition_hz,
        odor_labels=dataset.odor_labels,
        glomerulus_labels=dataset.glomerulus_labels,
        source_variables=tuple((*dataset.source_variables, source_name)),
    )


def _load_hdf5_mat_payload(data: bytes) -> dict[str, Any]:
    try:
        import h5py
    except ImportError as exc:
        raise ValueError("HDF5 MATLAB file requires h5py") from exc
    with h5py.File(BytesIO(data), "r") as handle:
        return {key: _hdf5_to_plain(value) for key, value in handle.items()}


def _hdf5_to_plain(value: Any) -> Any:
    if hasattr(value, "items"):
        return {key: _hdf5_to_plain(child) for key, child in value.items()}
    return np.asarray(value)


def _dataframe_summaries_from_workbook(
    dataset_id: str,
    workbook_name: str,
    data: bytes,
) -> list[dict[str, int | float | str]]:
    summaries: list[dict[str, int | float | str]] = []
    try:
        sheets = pd.read_excel(BytesIO(data), sheet_name=None, header=None)
    except Exception:
        return summaries
    for sheet_name, frame in sheets.items():
        frame = frame.dropna(how="all").dropna(axis=1, how="all")
        if frame.empty:
            continue
        numeric = frame.apply(pd.to_numeric, errors="coerce")
        numeric_values = numeric.to_numpy(dtype=float, copy=False)
        finite_numeric = np.isfinite(numeric_values)
        summaries.append(
            {
                "dataset_id": dataset_id,
                "workbook": workbook_name,
                "sheet": str(sheet_name),
                "row_count": int(frame.shape[0]),
                "column_count": int(frame.shape[1]),
                "numeric_cell_count": int(finite_numeric.sum()),
                "finite_numeric_fraction": float(finite_numeric.mean())
                if numeric_values.size
                else 0.0,
            }
        )
    return summaries


def _iter_xlsx_members(archive: zipfile.ZipFile, prefix: str = ""):
    for info in archive.infolist():
        if info.is_dir():
            continue
        member_name = f"{prefix}{info.filename}"
        data = archive.read(info)
        if info.filename.lower().endswith(".xlsx"):
            yield member_name, data
        elif info.filename.lower().endswith(".zip"):
            with zipfile.ZipFile(BytesIO(data)) as nested:
                yield from _iter_xlsx_members(nested, prefix=f"{info.filename}/")


def _panels_from_workbook(
    dataset_id: str, workbook_name: str, data: bytes
) -> list[EmpiricalOdorResponsePanel]:
    panels: list[EmpiricalOdorResponsePanel] = []
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    receptor = _receptor_from_name(workbook_name)
    for sheet in workbook.worksheets:
        rows = [tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)]
        table = _trim_empty_rows(rows)
        if len(table) < 2:
            continue
        first_header = str(table[0][0] or "").strip().lower()
        if first_header == "frames":
            rows_for_panel = _timecourse_rows(sheet.title, table, receptor)
            modality = f"{workbook_name}:{sheet.title}:timecourse_peak"
        elif first_header.startswith("bees"):
            rows_for_panel = _stimulus_by_bee_rows(table)
            modality = f"{workbook_name}:{sheet.title}:regional_intensity"
        elif first_header.startswith("individual"):
            rows_for_panel = _individual_by_stimulus_rows(table, receptor or sheet.title)
            modality = f"{workbook_name}:{sheet.title}:receptor_response"
        elif workbook_name.lower().startswith("data "):
            rows_for_panel = _nouvian_rows_from_sheet(workbook_name, sheet.title, table)
            modality = f"{workbook_name}:{sheet.title}:nouvian_structured_response"
        else:
            rows_for_panel = _generic_numeric_rows(table, receptor or sheet.title)
            modality = f"{workbook_name}:{sheet.title}:generic_numeric_response"
        if not rows_for_panel:
            continue
        panel = parse_tabular_odor_response_rows(
            dataset_id,
            rows_for_panel,
            stimulus_key="stimulus",
            response_key="response",
            channel_key="channel",
            modality=modality,
            units=_units_from_sheet(sheet.title),
        )
        panels.append(panel)
    workbook.close()
    return panels


def _trim_empty_rows(rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    return [row for row in rows if any(cell is not None for cell in row)]


def _stimulus_by_bee_rows(table: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    headers = [str(cell).strip() if cell is not None else "" for cell in table[0]]
    rows: list[dict[str, Any]] = []
    for row in table[1:]:
        stimulus = row[0]
        if stimulus in (None, ""):
            continue
        for index, channel in enumerate(headers[1:], start=1):
            if not channel:
                continue
            value = _numeric(row[index] if index < len(row) else None)
            if value is not None:
                rows.append({"stimulus": str(stimulus), "channel": channel, "response": value})
    return rows


def _individual_by_stimulus_rows(
    table: list[tuple[Any, ...]], channel_prefix: str
) -> list[dict[str, Any]]:
    headers = [str(cell).strip() if cell is not None else "" for cell in table[0]]
    rows: list[dict[str, Any]] = []
    for row in table[1:]:
        individual = str(row[0]).strip() if row and row[0] is not None else ""
        if not individual:
            continue
        for index, stimulus in enumerate(headers[1:], start=1):
            if not stimulus:
                continue
            value = _numeric(row[index] if index < len(row) else None)
            if value is not None:
                rows.append(
                    {
                        "stimulus": stimulus,
                        "channel": f"{channel_prefix}:{individual}",
                        "response": value,
                    }
                )
    return rows


def _timecourse_rows(
    sheet_name: str, table: list[tuple[Any, ...]], channel_prefix: str | None
) -> list[dict[str, Any]]:
    headers = [str(cell).strip() if cell is not None else "" for cell in table[0]]
    rows: list[dict[str, Any]] = []
    for index, stimulus in enumerate(headers[1:], start=1):
        values = [
            abs(value)
            for row in table[1:]
            if (value := _numeric(row[index] if index < len(row) else None)) is not None
        ]
        if values:
            rows.append(
                {
                    "stimulus": stimulus,
                    "channel": f"{channel_prefix or 'timecourse'}:{sheet_name}",
                    "response": max(values),
                }
            )
    return rows


def _nouvian_rows_from_sheet(
    workbook_name: str,
    sheet_name: str,
    table: list[tuple[Any, ...]],
) -> list[dict[str, Any]]:
    lowered = f"{workbook_name}:{sheet_name}".lower()
    if "hplc" in lowered:
        return _nouvian_hplc_rows(table)
    return _nouvian_binary_behaviour_rows(sheet_name, table)


def _nouvian_hplc_rows(table: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    headers = [str(cell).strip() if cell is not None else "" for cell in table[0]]
    normalized = [_normalize_header(header) for header in headers]
    if "brain_region" not in normalized:
        return []
    region_idx = normalized.index("brain_region")
    colony_idx = normalized.index("colony") if "colony" in normalized else None
    odor_idx = normalized.index("odour") if "odour" in normalized else None
    behavior_idx = normalized.index("behaviour") if "behaviour" in normalized else None
    amine_columns = [
        (idx, _amine_label(header))
        for idx, header in enumerate(headers)
        if _amine_label(header) is not None
    ]
    rows: list[dict[str, Any]] = []
    for source_row in table[1:]:
        region = _text(source_row[region_idx] if region_idx < len(source_row) else None)
        if not region or len(region) > 12:
            continue
        stimulus = _join_label(
            region,
            _text(
                source_row[odor_idx]
                if odor_idx is not None and odor_idx < len(source_row)
                else None
            ),
            _text(
                source_row[behavior_idx]
                if behavior_idx is not None and behavior_idx < len(source_row)
                else None
            ),
        )
        colony = _text(
            source_row[colony_idx]
            if colony_idx is not None and colony_idx < len(source_row)
            else None
        )
        for idx, amine in amine_columns:
            value = _numeric(source_row[idx] if idx < len(source_row) else None)
            if value is not None:
                rows.append(
                    {
                        "stimulus": stimulus,
                        "channel": _join_label(amine or "amine", colony),
                        "response": value,
                    }
                )
    return rows


def _nouvian_binary_behaviour_rows(
    sheet_name: str, table: list[tuple[Any, ...]]
) -> list[dict[str, Any]]:
    if len(table) < 3:
        return []
    header_idx = _binary_header_index(table)
    if header_idx is None:
        return []
    header = table[header_idx]
    super_header = table[header_idx - 1] if header_idx > 0 else tuple()
    filled_super = _fill_forward(tuple(_text(cell) for cell in super_header))
    stimuli = {
        idx: _text(cell)
        for idx, cell in enumerate(header)
        if idx > 0
        and _text(cell)
        and any(
            _numeric(row[idx] if idx < len(row) else None) is not None
            for row in table[header_idx + 1 :]
        )
    }
    rows: list[dict[str, Any]] = []
    current_row_channel = ""
    for row_index, row in enumerate(table[header_idx + 1 :], start=1):
        label = _text(row[0] if row else None)
        if rows and _is_nouvian_stats_label(label):
            break
        if rows and not any(
            _numeric(row[idx] if idx < len(row) else None) is not None for idx in stimuli
        ):
            break
        if label:
            current_row_channel = label
        for idx, stimulus in stimuli.items():
            value = _numeric(row[idx] if idx < len(row) else None)
            if value is None:
                continue
            column_channel = filled_super[idx] if idx < len(filled_super) else ""
            if column_channel and column_channel.lower() != "odour and treatment":
                channel = _join_label(sheet_name, column_channel)
            else:
                channel = _join_label(sheet_name, current_row_channel or f"replicate_{row_index}")
            rows.append({"stimulus": stimulus, "channel": channel, "response": value})
    return rows


def _binary_header_index(table: list[tuple[Any, ...]]) -> int | None:
    candidates: list[tuple[int, int, int]] = []
    for idx, row in enumerate(table[:4]):
        text_cells = [_text(cell) for cell in row[1:]]
        if sum(bool(cell) for cell in text_cells) >= 2 and any(
            _numeric(next_row[col] if col < len(next_row) else None) is not None
            for next_row in table[idx + 1 : idx + 5]
            for col in range(1, min(len(row), len(next_row)))
        ):
            odorish = sum(_is_odor_condition_label(cell) for cell in text_cells)
            descriptor = int(any(token in _text(row[0]).lower() for token in ("odour", "odor")))
            candidates.append((odorish, descriptor, idx))
    if not candidates:
        return None
    return max(candidates)[2]


def _fill_forward(values: tuple[str, ...]) -> tuple[str, ...]:
    out: list[str] = []
    current = ""
    for value in values:
        current = value or current
        out.append(current)
    return tuple(out)


def _normalize_header(header: str) -> str:
    return " ".join(header.lower().replace("(", " ").replace(")", " ").split()).replace(" ", "_")


def _amine_label(header: str) -> str | None:
    normalized = _normalize_header(header)
    if "octopamine" in normalized:
        return "octopamine"
    if "dopamine" in normalized:
        return "dopamine"
    if "serotonin" in normalized:
        return "serotonin"
    return None


def _is_odor_condition_label(text: str) -> bool:
    lowered = text.lower()
    return any(token in lowered for token in ("tec", "iaa"))


def _is_nouvian_stats_label(text: str) -> bool:
    lowered = text.lower()
    return lowered in {
        "mean",
        "<none>",
        "odour",
        "odor",
        "treatment",
        "odour*treatment",
        "odor*treatment",
        "(intercept)",
    } or lowered.startswith("glm")


def _generic_numeric_rows(
    table: list[tuple[Any, ...]], channel_prefix: str
) -> list[dict[str, Any]]:
    headers = [str(cell).strip() if cell is not None else "" for cell in table[0]]
    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(table[1:], start=1):
        stimulus = _first_text_value(row) or f"row_{row_index}"
        for index, header in enumerate(headers):
            value = _numeric(row[index] if index < len(row) else None)
            if value is not None:
                channel = header or f"column_{index + 1}"
                rows.append(
                    {
                        "stimulus": stimulus,
                        "channel": f"{channel_prefix}:{channel}",
                        "response": value,
                    }
                )
    stimulus_count = len({row["stimulus"] for row in rows})
    channel_count = len({row["channel"] for row in rows})
    return rows if stimulus_count >= 2 and channel_count >= 1 else []


def _first_text_value(row: tuple[Any, ...]) -> str | None:
    for value in row:
        if value in (None, ""):
            continue
        if _numeric(value) is None:
            text = str(value).strip()
            if text:
                return text
    return None


def _join_label(*parts: str | None) -> str:
    return ":".join(part for part in parts if part)


def _text(value: Any) -> str:
    return "" if value in (None, "") else str(value).strip()


def _numeric(value: Any) -> float | None:
    if value in (None, "", "NA"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _receptor_from_name(name: str) -> str | None:
    lowered = name.lower()
    if "amor109" in lowered:
        return "AmOR109"
    if "amor136" in lowered:
        return "AmOR136"
    return None


def _units_from_sheet(sheet_name: str) -> str:
    lowered = sheet_name.lower()
    if "hplc" in lowered:
        return "pg/ul"
    if lowered in {"behaviour", "5ht", "da", "cyp", "flu"}:
        return "attack_or_stinging_probability"
    if "spike" in lowered:
        return "spikes/s"
    if "distance" in lowered:
        return "euclidean distance"
    return "normalized fluorescence"


def main() -> None:
    cfg = load_config()
    source_dir = PROJECT_ROOT / "output" / "data" / "empirical_sources"
    panels = load_empirical_panels(source_dir)
    workbook_summaries = workbook_dataframe_summaries(source_dir)
    antennal_summaries = load_antennal_movement_summaries(source_dir)
    waggle_dataset = load_waggle_follower_dataset(source_dir)
    calcium_datasets = load_calcium_datasets(source_dir)
    anatomy_summary, anatomy_inventories, neuropil_abbreviations, anatomy_records = load_anatomy(
        source_dir
    )
    if not panels:
        raise SystemExit(
            "No empirical panels found. Run scripts/fetch_empirical_bee_data.py first."
        )
    stats = tuple(analyze_odor_panel(panel) for panel in panels)
    calcium_summaries = tuple(
        summarize_calcium_trials(dataset.traces, dataset.protocol(cfg))
        for dataset in calcium_datasets
    )
    bank = _integrated_template_bank(panels, calcium_datasets, cfg).limited(24)
    drive_label = _select_drive_label(bank.templates)
    empirical_drive = bank.templates[drive_label]
    empirical_vibration = (
        antennal_vibration_from_movement(antennal_summaries[0]) if antennal_summaries else None
    )
    if waggle_dataset is not None:
        empirical_vibration = waggle_vibration_from_followers(waggle_dataset.summary)
    observation = zero_observation(cfg)
    observation_updates: dict[str, Any] = {"antennal_channels": empirical_drive}
    if empirical_vibration is not None:
        observation_updates["antennal_vibration"] = empirical_vibration
    observation = observation.__class__(**{**observation.__dict__, **observation_updates})
    brain = process_observation(observation, cfg, seed=cfg.seed, odor_templates=bank.templates)
    result = run_simulation(
        cfg,
        steps=24,
        empirical_odor_templates=bank.templates,
        empirical_drive=empirical_drive,
        empirical_antennal_vibration=empirical_vibration,
    )
    alignment = dict(
        sorted(brain.empirical_alignment.items(), key=lambda item: item[1], reverse=True)
    )
    quality_checks = (
        tuple(check for row in stats for check in row.quality_checks)
        + bank.quality_checks
        + tuple(
            check for dataset in calcium_datasets for check in validate_calcium_dataset(dataset)
        )
    )
    activity_summary = activity_summary_from_components(
        panels,
        stats,
        bank,
        antennal_summaries=antennal_summaries,
        calcium_summaries=calcium_summaries,
        waggle_follower_summary=waggle_dataset.summary if waggle_dataset is not None else None,
        quality_checks=quality_checks,
    )
    archive_status = _archive_status(source_dir)
    anatomy_download_status = _anatomy_download_status(source_dir)
    downloaded_dataset_ids = _downloaded_dataset_ids(
        archive_status,
        _catalog_status(source_dir),
        anatomy_download_status,
    )
    parseable_dataset_ids = _parseable_dataset_ids(
        panels, calcium_datasets, antennal_summaries, waggle_dataset, anatomy_summary
    )
    data_completeness = bee_brain_data_completeness_panel(
        empirical_brain_datasets(),
        downloaded_dataset_ids=tuple(downloaded_dataset_ids),
        parseable_dataset_ids=tuple(parseable_dataset_ids),
        parser_status_by_dataset=_parser_status_by_dataset(
            _catalog_status(source_dir), parseable_dataset_ids
        ),
        parseability_target=0.8,
    )
    known_gaps = _known_gaps(archive_status, anatomy_summary, activity_summary, waggle_dataset)
    report = {
        "panel_count": len(panels),
        "workbook_dataframe_summary_count": len(workbook_summaries),
        "calcium_dataset_count": len(calcium_datasets),
        "antennal_movement_summary_count": len(antennal_summaries),
        "waggle_follower_dataset_available": waggle_dataset is not None,
        "template_count": len(bank.templates),
        "all_quality_checks_passed": all_checks_passed(quality_checks),
        "activity_summary": activity_summary.as_dict(),
        "anatomy_summary": anatomy_summary.as_dict(),
        "anatomy_inventories": [inventory.as_dict() for inventory in anatomy_inventories],
        "neuropil_abbreviations": [entry.as_dict() for entry in neuropil_abbreviations],
        "calcium_datasets": [dataset.as_summary_dict() for dataset in calcium_datasets],
        "calcium_summaries": [summary.as_dict() for summary in calcium_summaries],
        "panels": [panel.as_summary_dict() for panel in panels],
        "antennal_movement_summaries": [summary.as_dict() for summary in antennal_summaries],
        "waggle_follower_analysis": waggle_dataset.as_dict() if waggle_dataset is not None else {},
        "panel_stats": [row.as_dict() for row in stats],
        "workbook_dataframe_summaries": workbook_summaries,
        "brain_data_completeness": data_completeness.as_dict(),
        "template_sources": bank.sources,
        "empirical_drive_label": drive_label,
        "empirical_antennal_vibration": empirical_vibration.tolist()
        if empirical_vibration is not None
        else None,
        "archive_status": archive_status,
        "anatomy_downloads": anatomy_download_status,
        "template_alignment_matrix": template_alignment_matrix(bank),
        "stack_alignment": alignment,
        "simulation_summary_with_empirical_templates": result.summary(),
        "known_gaps": known_gaps,
    }
    data_dir = PROJECT_ROOT / "output" / "data"
    figure_paths = generate_empirical_figures(
        panels,
        stats,
        alignment,
        PROJECT_ROOT / "output" / "figures" / "empirical",
        antennal_summaries=antennal_summaries,
        anatomy_summary=anatomy_summary,
        anatomy_inventories=anatomy_inventories,
        neuropil_abbreviations=neuropil_abbreviations,
        activity_summary=activity_summary,
        waggle_summary=waggle_dataset.summary if waggle_dataset is not None else None,
        data_completeness=data_completeness,
    )
    end_to_end_report = BeeBrainEndToEndReport(
        anatomy=anatomy_summary,
        activity=activity_summary,
        dataset_ids=tuple(sorted({panel.dataset_id for panel in panels})),
        figure_paths=tuple(str(path) for path in figure_paths),
        archive_status=tuple(archive_status),
        anatomy_downloads=tuple(record.as_dict() for record in anatomy_records),
        known_gaps=tuple(known_gaps),
    )
    report["end_to_end_report"] = end_to_end_report.as_dict()
    write_json(data_dir / "empirical_analysis.json", report)
    write_json(data_dir / "empirical_template_bank.json", bank.as_dict())
    write_json(data_dir / "bee_brain_end_to_end_report.json", end_to_end_report.as_dict())
    write_json(
        data_dir / "waggle_follower_analysis.json",
        waggle_dataset.as_dict() if waggle_dataset is not None else {"dataset_id": "missing"},
    )
    write_json(data_dir / "brain_data_completeness.json", data_completeness.as_dict())
    markdown = _markdown_report(report, figure_paths)
    (PROJECT_ROOT / "output" / "reports" / "empirical_analysis.md").write_text(
        markdown, encoding="utf-8"
    )
    (PROJECT_ROOT / "output" / "reports" / "waggle_follower_analysis.md").write_text(
        _waggle_markdown(waggle_dataset, data_completeness), encoding="utf-8"
    )
    write_signposts(PROJECT_ROOT)
    write_project_readiness_review(PROJECT_ROOT)
    print(
        f"Analyzed {len(panels)} empirical panels, {len(calcium_datasets)} calcium datasets, "
        f"{len(antennal_summaries)} antennal summaries, {len(anatomy_inventories)} anatomy "
        f"inventories, and integrated {len(bank.templates)} templates"
    )


def _integrated_template_bank(
    panels: tuple[EmpiricalOdorResponsePanel, ...],
    calcium_datasets: tuple[EmpiricalCalciumDataset, ...],
    cfg: BeeStackConfig,
) -> EmpiricalTemplateBank:
    bank = build_empirical_template_bank(panels, cfg)
    templates = dict(bank.templates)
    sources = {label: tuple(source) for label, source in bank.sources.items()}
    checks = list(bank.quality_checks)
    for dataset in calcium_datasets:
        checks.extend(validate_calcium_dataset(dataset))
        for label, vector in response_templates_from_calcium_dataset(dataset, cfg).items():
            key = f"calcium {_normalize_label(label)}"
            templates[key] = vector
            sources[key] = (f"{dataset.dataset_id}:matlab_calcium",)
    return EmpiricalTemplateBank(templates, sources, tuple(checks))


def _normalize_label(label: str) -> str:
    return " ".join(str(label).replace("_", " ").split()).strip().lower()


def _archive_status(source_dir: Path) -> list[dict[str, Any]]:
    path = source_dir / "archives.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def _catalog_status(source_dir: Path) -> list[dict[str, Any]]:
    return _json_rows(source_dir / "catalog.json")


def _anatomy_download_status(source_dir: Path) -> list[dict[str, Any]]:
    return _json_rows(source_dir / "anatomy_downloads.json")


def _json_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    return [dict(row) for row in payload] if isinstance(payload, list) else []


def _downloaded_dataset_ids(
    archive_status: list[dict[str, Any]],
    catalog_status: list[dict[str, Any]],
    anatomy_status: list[dict[str, Any]],
) -> set[str]:
    ids = {
        str(row.get("dataset_id"))
        for row in archive_status
        if row.get("archive_downloaded") or int(row.get("file_downloaded_count") or 0) > 0
    }
    ids.update(str(row.get("dataset_id")) for row in catalog_status if row.get("file_downloaded"))
    ids.update(str(row.get("dataset_id")) for row in anatomy_status if row.get("downloaded"))
    return {dataset_id for dataset_id in ids if dataset_id and dataset_id != "None"}


def _parseable_dataset_ids(
    panels: tuple[EmpiricalOdorResponsePanel, ...],
    calcium_datasets: tuple[EmpiricalCalciumDataset, ...],
    antennal_summaries: tuple[AntennalMovementSummary, ...],
    waggle_dataset: EmpiricalWaggleFollowerDataset | None,
    anatomy_summary: BeeBrainAnatomySummary,
) -> set[str]:
    ids = {panel.dataset_id for panel in panels}
    ids.update(dataset.dataset_id for dataset in calcium_datasets)
    ids.update(summary.dataset_id for summary in antennal_summaries)
    if waggle_dataset is not None:
        ids.add(waggle_dataset.dataset_id)
    if anatomy_summary.inventory_count or anatomy_summary.neuropil_count:
        ids.add(anatomy_summary.dataset_id)
    return ids


def _parser_status_by_dataset(
    catalog_status: list[dict[str, Any]],
    parseable_dataset_ids: set[str],
) -> dict[str, str]:
    statuses: dict[str, str] = {dataset_id: "parsed" for dataset_id in parseable_dataset_ids}
    for row in catalog_status:
        dataset_id = str(row.get("dataset_id") or "")
        if not dataset_id or statuses.get(dataset_id) == "parsed":
            continue
        parser_status = str(row.get("parser_status") or "")
        file_error = str(row.get("file_error") or "")
        if file_error:
            statuses[dataset_id] = f"download_blocked:{file_error}"
        elif parser_status and _parser_status_priority(parser_status) > _parser_status_priority(
            statuses.get(dataset_id, "")
        ):
            statuses[dataset_id] = parser_status
    return statuses


def _parser_status_priority(status: str) -> int:
    if status.startswith("download_blocked:"):
        return 100
    if status.startswith("supported_"):
        return 80
    if status.startswith("metadata_"):
        return 20
    if status:
        return 10
    return 0


def _known_gaps(
    archive_status: list[dict[str, Any]],
    anatomy_summary: BeeBrainAnatomySummary,
    activity_summary: Any,
    waggle_dataset: EmpiricalWaggleFollowerDataset | None = None,
) -> list[str]:
    gaps: list[str] = []
    for row in archive_status:
        if not row.get("archive_downloaded") and int(row.get("file_downloaded_count") or 0) == 0:
            gaps.append(
                f"{row.get('dataset_id')} cataloged but no archive or file-level payload is local"
            )
    if anatomy_summary.downloaded_asset_count < anatomy_summary.asset_count:
        gaps.append("some Honeybee Standard Brain anatomy assets are cataloged but not local")
    if activity_summary.calcium_dataset_count == 0:
        gaps.append("Paoli MATLAB calcium traces are not yet local or parseable")
    if anatomy_summary.neuropil_count == 0:
        gaps.append("neuropil abbreviation HTML is not yet local or parseable")
    if waggle_dataset is None:
        gaps.append("Hadjitofi-Webb Figshare waggle-following CSVs are not yet local or parseable")
    return gaps


def _markdown_report(report: dict[str, Any], figure_paths: list[Path]) -> str:
    top = list(report["stack_alignment"].items())[:8]
    lines = [
        "# Empirical Bee Data Analysis",
        "",
        f"- Panels analyzed: `{report['panel_count']}`",
        f"- Workbook dataframe summaries: `{report['workbook_dataframe_summary_count']}`",
        f"- Calcium datasets parsed: `{report['calcium_dataset_count']}`",
        f"- Antennal movement summaries: `{report['antennal_movement_summary_count']}`",
        f"- Anatomy inventories parsed: `{report['anatomy_summary']['inventory_count']}`",
        f"- Neuropil abbreviations parsed: `{report['anatomy_summary']['neuropil_count']}`",
        f"- Templates integrated into BeeBrain: `{report['template_count']}`",
        f"- All quality checks passed: `{report['all_quality_checks_passed']}`",
        f"- Mean odor separability: `{report['activity_summary']['mean_odor_separability']:.3f}`",
        f"- Empirical drive label: `{report['empirical_drive_label']}`",
        f"- Empirical antennal vibration: `{report['empirical_antennal_vibration']}`",
        f"- Final stack empirical odor: `{report['simulation_summary_with_empirical_templates']['final_empirical_odor']}`",
        f"- Final stack empirical alignment: `{report['simulation_summary_with_empirical_templates']['final_empirical_alignment']:.3f}`",
        "",
        "## Top BeeBrain Alignments",
        "",
    ]
    for label, value in top:
        lines.append(f"- `{label}`: {value:.3f}")
    lines.extend(["", "## Anatomy Coverage", ""])
    lines.append(
        "- Downloaded assets: `{downloaded_asset_count}` of `{asset_count}`; "
        "VRML files `{vrml_file_count}`; TIFF files `{tiff_file_count}`; "
        "uncompressed bytes `{total_uncompressed_bytes}`.".format(**report["anatomy_summary"])
    )
    region_counts = report["anatomy_summary"]["major_regions"]
    if region_counts:
        for region, count in sorted(region_counts.items()):
            lines.append(f"- `{region}` abbreviations: `{count}`")
    else:
        lines.append("- No neuropil abbreviation table was available locally.")
    lines.extend(["", "## Activity Summary", ""])
    activity = report["activity_summary"]
    lines.append(
        f"- Calcium latency `{activity['calcium_mean_latency_s']:.3f}` s; "
        f"inhibitory fraction `{activity['calcium_inhibitory_fraction']:.3f}`; "
        f"excitatory fraction `{activity['calcium_excitatory_fraction']:.3f}`."
    )
    lines.append(
        f"- Aftersmell/post-odor response mean: `{activity['aftersmell_response_mean']:.3f}`"
    )
    for region, value in sorted(activity["region_response_means"].items()):
        lines.append(f"- `{region}` mean absolute response: `{value:.3f}`")
    lines.extend(["", "## Antennal Active Sensing", ""])
    for summary in report["antennal_movement_summaries"]:
        lines.append(
            "- `{dataset_id}`: rows `{row_count}`, bees `{bee_count}`, plumes `{plume_count}`, "
            "odor-on `{odor_on_fraction:.3f}`, left/right correlation `{left_right_theta_correlation:.3f}`".format(
                **summary
            )
        )
    if not report["antennal_movement_summaries"]:
        lines.append("- No frame-level antennal summaries were available in local archives.")
    waggle = report.get("waggle_follower_analysis") or {}
    if waggle:
        summary = waggle.get("summary", {})
        lines.extend(["", "## Waggle Follower Decoding", ""])
        lines.append(
            "- Hadjitofi-Webb follower tracks `{track_count}`; confidence `{confidence_score:.3f}`; "
            "angle/midpoint correlation `{follower_angle_midpoint_correlation:.3f}`; "
            "decoding improvement `{decoding_improvement_fraction:.3f}`.".format(**summary)
        )
    completeness = report.get("brain_data_completeness", {})
    if completeness:
        lines.extend(["", "## Brain Data Completeness", ""])
        lines.append(
            "- Curated datasets `{dataset_count}`; downloaded `{downloaded_dataset_count}`; "
            "parseable `{parseable_dataset_count}`; parseable fraction `{parseable_fraction:.3f}`; "
            "source-verified `{source_verified_dataset_count}`; target satisfied `{parseability_target_satisfied}`.".format(
                **completeness
            )
        )
    blocked_archives = [
        row for row in report["archive_status"] if not row.get("archive_downloaded")
    ]
    if blocked_archives:
        lines.extend(["", "## Archive Caveats", ""])
        for row in blocked_archives:
            lines.append(
                f"- `{row['dataset_id']}` was cataloged but not downloaded"
                f" ({row.get('archive_error') or 'file-level fallback did not produce a local file'})."
            )
    if report["known_gaps"]:
        lines.extend(["", "## Known Gaps", ""])
        for gap in report["known_gaps"]:
            lines.append(f"- {gap}.")
    lines.extend(["", "## Figures", ""])
    for path in figure_paths:
        lines.append(f"- `{path}`")
    lines.append("")
    return "\n".join(lines)


def _waggle_markdown(
    waggle_dataset: EmpiricalWaggleFollowerDataset | None,
    data_completeness: Any,
) -> str:
    lines = [
        "# Waggle Follower Analysis",
        "",
        "This report covers the curated Hadjitofi-Webb Figshare waggle-following source "
        "and its integration into BeeBrain/BeeSwarm waggle decoding.",
        "",
    ]
    if waggle_dataset is None:
        lines.extend(
            [
                "- Dataset status: `not local or not parseable`",
                "- Expected source: `https://figshare.com/articles/dataset/Honeybee_antennal_positioning_data_when_following_dances/24715977`",
                "- Regeneration command: `uv run python scripts/fetch_empirical_bee_data.py`",
                "",
            ]
        )
    else:
        summary = waggle_dataset.summary
        lines.extend(
            [
                "- Dataset status: `parsed`",
                f"- Tracks: `{summary.track_count}`",
                f"- Feature rows: `{summary.feature_row_count}`",
                f"- Binned rows: `{summary.binned_row_count}`",
                f"- Model-error rows: `{summary.model_error_row_count}`",
                f"- Follower angle/midpoint correlation: `{summary.follower_angle_midpoint_correlation:.3f}`",
                f"- Both-antennae decoding error: `{summary.both_antennae_mean_abs_error_deg:.3f}` deg",
                f"- No-antennae decoding error: `{summary.no_antennae_mean_abs_error_deg:.3f}` deg",
                f"- Decoding improvement fraction: `{summary.decoding_improvement_fraction:.3f}`",
                f"- Confidence score: `{summary.confidence_score:.3f}`",
                "",
                "## Source Files",
                "",
            ]
        )
        lines.extend(f"- `{path}`" for path in waggle_dataset.source_files[:24])
        lines.append("")
    completeness = data_completeness.as_dict()
    lines.extend(
        [
            "## BeeBrain Data Completeness",
            "",
            f"- Curated datasets: `{completeness['dataset_count']}`",
            f"- Downloaded datasets: `{completeness['downloaded_dataset_count']}`",
            f"- Parseable datasets: `{completeness['parseable_dataset_count']}`",
            f"- Parseable fraction: `{completeness['parseable_fraction']:.3f}`",
            f"- Source-verified datasets: `{completeness['source_verified_dataset_count']}`",
            f"- Source-verified blockers: `{completeness['source_verified_blocked_count']}`",
            f"- Parseability target satisfied: `{completeness['parseability_target_satisfied']}`",
            "",
        ]
    )
    return "\n".join(lines)


def _select_drive_label(templates: dict[str, Any]) -> str:
    preferred = ("alarm", "isopentyl acetate", "16ol", "1-octanol")
    for label in preferred:
        if label in templates:
            return label
    return sorted(templates)[0]


def _dataset_id_from_path(source_dir: Path, path: Path) -> str:
    relative = path.relative_to(source_dir)
    return relative.parts[0]


if __name__ == "__main__":
    main()
